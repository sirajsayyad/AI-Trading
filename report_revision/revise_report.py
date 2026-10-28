from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path('E:/AI_Trading/report_revision')
SOURCE = Path('C:/Users/siraj/Downloads/Mini project Phase 1 report SS - Updated.docx')
OUT = Path('C:/Users/siraj/Downloads/Mini project Phase 1 report SS - Final.docx')
src = Document(SOURCE)
def clean(t):
    t = re.sub(r'(\d)\ufffd(?=\d)', r'\1-', t)
    t = t.replace('SVERI\ufffdS', "SVERI'S").replace('Flor\ufffdncio', 'Florêncio')
    t = re.sub(r'\ufffd(?=s\b)', "'", t)
    t = t.replace('Kaliski\ufffds', "Kaliski's").replace('users\ufffd', "users'")
    t = t.replace(' \ufffd ', ' - ').replace('\ufffd', '"')
    return ' '.join(t.split())

with ZipFile(SOURCE) as z:
    (ROOT/'crest.jpg').write_bytes(z.read('word/media/image1.jpeg'))

# A clean rendering of the architecture already described in the source report.
im = Image.new('RGB', (1600, 1020), 'white')
dr = ImageDraw.Draw(im)
font = ImageFont.truetype('C:/Windows/Fonts/times.ttf', 32)
small = ImageFont.truetype('C:/Windows/Fonts/times.ttf', 27)
def box(rect, lines):
    dr.rounded_rectangle(rect, radius=12, outline='black', width=3, fill='#f3f3f3')
    for j,line in enumerate(lines):
        dr.text(((rect[0]+rect[2])/2, (rect[1]+rect[3])/2+(j-(len(lines)-1)/2)*38), line, font=font, anchor='mm', fill='black')
def arrow(a,b,label=None):
    dr.line([a,b],fill='black',width=4)
    import math
    ang=math.atan2(b[1]-a[1],b[0]-a[0]); size=16
    pts=[b,(b[0]-size*math.cos(ang-.5),b[1]-size*math.sin(ang-.5)),(b[0]-size*math.cos(ang+.5),b[1]-size*math.sin(ang+.5))]
    dr.polygon(pts,fill='black')
    if label: dr.text(((a[0]+b[0])/2+18,(a[1]+b[1])/2),label,font=small,fill='black',anchor='lm')
box((560,20,1040,115),['User and application interface'])
box((560,185,1040,290),['Master-password authentication'])
box((560,360,1040,470),['Key derivation','Salt and configured work factor'])
box((560,540,1040,650),['Encryption and decryption','Authenticated encryption'])
box((560,750,1040,870),['Credential vault','Ciphertext, nonce and tag'])
box((35,360,465,470),['Credential management','Add, view, edit, delete, search'])
box((1135,360,1565,470),['Password utilities','Generator and strength checker'])
box((35,750,465,870),['Session security','Timeout, lock and logout'])
for a,b in [((800,115),(800,185)),((800,290),(800,360)),((800,470),(800,540)),((800,650),(800,750)),((465,415),(560,415)),((1135,415),(1040,415))]: arrow(a,b)
dr.line([(250,750),(250,590),(560,590)],fill='black',width=3)
arrow((465,590),(560,590))
dr.text((800,945),'Decrypted credentials are available only within an authorized session.',font=font,anchor='mm',fill='black')
im.save(ROOT/'architecture.png')

d=Document()
normal=d.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.35
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name,size in [('Title',20),('Heading 1',16),('Heading 2',13),('Heading 3',12),('Caption',11)]:
    s=d.styles[name]; s.font.name='Times New Roman'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s.font.bold=name!='Caption'; s.font.italic=False
    s.paragraph_format.space_before=Pt(10 if 'Heading' in name else 4)
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.keep_together=True
    s.paragraph_format.line_spacing=1.15
for name in ['List Bullet','List Number']:
    d.styles[name].font.name='Times New Roman'; d.styles[name].font.size=Pt(12)
    d.styles[name].paragraph_format.line_spacing=1.3
    d.styles[name].paragraph_format.space_after=Pt(5)

def setup(sec):
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=Inches(.85); sec.bottom_margin=Inches(.85)
    sec.left_margin=Inches(1.15); sec.right_margin=Inches(.85)
    sec.footer_distance=Inches(.4); sec.header_distance=Inches(.4)
setup(d.sections[0])
def para(t,style=None):
    p=d.add_paragraph(clean(t),style); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if not style else None
    return p
def center(t,size=12,bold=False,before=0,after=8):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=1.25
    r=p.add_run(t); r.bold=bold; r.font.size=Pt(size)
    return p
def head(t,level=2): return d.add_heading(t,level)
def numbering(sec,fmt,start=None):
    for old in sec._sectPr.findall(qn('w:pgNumType')): sec._sectPr.remove(old)
    n=OxmlElement('w:pgNumType'); n.set(qn('w:fmt'),fmt)
    if start is not None: n.set(qn('w:start'),str(start))
    sec._sectPr.append(n)
    sec.footer.is_linked_to_previous=False
    p=sec.footer.paragraphs[0]; p.clear(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); r._r.addnext(f)
def section(fmt=None,start=None):
    s=d.add_section(WD_SECTION_START.NEW_PAGE); setup(s)
    if fmt: numbering(s,fmt,start)
    return s

center('Mini Project Phase I Report',18,True,after=16)
center('Secure Password Manager with\nMaster-Key Encryption',20,True,after=18)
center('Submitted in partial fulfillment of the requirements for the award of',11,after=6)
center('Final Year B.Tech. - Semester I\nHonors in Cyber Security',13,True,after=8)
center('CSE Hon-05B Mini Project',12,True,after=12)
center('Submitted to',11,after=4)
center('Punyashlok Ahilyadevi Holkar Solapur University, Solapur',12,True,after=16)
center('Submitted by',12,True,after=6)
for t in ['Mr. Siraj Riyaj Sayyad | Roll No. 70 | LY B','Mr. Prathamesh Jayantkumar Sathe | Roll No. 68 | LY B','Mr. Adesh Dattatray Dethe | Roll No. 46 | LY B']:
    center(t,12,after=6)
center('Project Group No. B2',12,True,after=14)
center('Under the guidance of\nMs. A. S. Pathan',12,True,after=12)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ROOT/'crest.jpg'),width=Inches(1.0))
center('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING',12,True,after=6)
center("SVERI'S COLLEGE OF ENGINEERING, PANDHARPUR",13,True,after=10)
center('Academic Year 2026-2027',12,True)

section()
center("SVERI'S COLLEGE OF ENGINEERING, PANDHARPUR",15,True,after=10)
center('(Affiliated to Punyashlok Ahilyadevi Holkar Solapur University, Solapur)',11,after=14)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ROOT/'crest.jpg'),width=Inches(1.0))
center('CERTIFICATE',18,True,before=14,after=22)
para('This is to certify that the Mini Project Phase I Report entitled')
center('Secure Password Manager with Master-Key Encryption',15,True,before=10,after=16)
para('has been successfully submitted by the following students of Final Year B.Tech., Class and Division LY B, Project Group B2:')
for t in ['Mr. Siraj Riyaj Sayyad - Roll No. 70','Mr. Prathamesh Jayantkumar Sathe - Roll No. 68','Mr. Adesh Dattatray Dethe - Roll No. 46']: center(t,12,True,after=10)
para('in partial fulfillment of the requirements for Final Year B.Tech. - Semester I of Honors in Cyber Security for the subject CSE Hon-05B Mini Project, as per the curriculum prescribed by Punyashlok Ahilyadevi Holkar Solapur University, Solapur, during the Academic Year 2026-2027.')
p=para('Ms. A. S. Pathan\t\tDr. Vanita Jadhav'); p.paragraph_format.space_before=Pt(42); p.runs[0].bold=True
p.text='Ms. A. S. Pathan\tDr. Vanita Jadhav'; p.runs[0].bold=True
p.paragraph_format.tab_stops.add_tab_stop(Inches(3.65))
p=para('Project Guide'); p.add_run('\tHead of Department CSE')
p.paragraph_format.tab_stops.add_tab_stop(Inches(3.65))
center('Dr. B. P. Ronge',12,True,before=36,after=4)
center('Principal',12,after=16)
p=para('Place: Pandharpur'); p.add_run('\tDate: __________________')
p.paragraph_format.tab_stops.add_tab_stop(Inches(3.65))

section('lowerRoman',1)
head('Abstract',1)
para(src.paragraphs[134].text)
head('Project Context')
para('The project focuses on the relationship between authentication and protected storage. Authentication establishes whether a user may open the vault; encryption protects recoverable credentials when the database is closed. A master password is therefore an access secret rather than a plaintext value saved beside the records. The proposed design uses a salt and a configurable derivation process to obtain the material needed for cryptographic operations.')
head('Scope and Evaluation')
para('Phase I defines the requirements, architecture, implementation plan, and evaluation approach for a local credential-management application. Development and testing will use synthetic accounts. Functional checks will cover record creation, retrieval, editing, deletion, and search. Security checks will examine incorrect-password handling, encrypted persistence, nonce handling, timeout, and logout. These are planned acceptance checks; the report does not present completed implementation results or measured performance values.')
head('Report Organization')
para('Chapter 1 introduces the project and its motivation. Chapter 2 reviews the literature and identifies the research gap. Chapter 3 defines the problem, objectives, and scope. Chapter 4 presents the proposed workflow, architecture, and requirements. Chapter 5 gives the activity plan and estimated expenditure. Chapter 6 describes expected results, limitations, and future enhancements. Chapter 7 concludes the Phase I proposal.')
para('Keywords: password manager; master password; key derivation; encrypted vault; authentication; session security.')

d.add_page_break(); p=para('Contents'); p.runs[0].bold=True; p.runs[0].font.size=Pt(16)
p=d.add_paragraph(); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u'); p._p.append(f)

chapter_positions=[]
def copy_range(a,b,skip=(),joins=None):
    joins=joins or {}
    for i in range(a,b):
        if i in skip: continue
        s=src.paragraphs[i]; t=clean(s.text)
        if not t or t.isdigit() or t in ['storage','ates']: continue
        if i in joins: t=clean(joins[i])
        if s.style.name.startswith('Heading'):
            level=int(s.style.name[-1]); head(t,level)
        else: para(t,'List Bullet' if s.style.name=='List Paragraph' else None)
def chapter(a,b,num):
    section('decimal',1 if num==1 else None)
    chapter_positions.append((num,len(d.paragraphs)))
    copy_range(a,b)

chapter(130,133,1)
copy_range(135,169)

chapter(169,218,2)

chapter(218,237,3)
head('Acceptance Boundary')
para('The demonstration will cover vault creation, authentication, synthetic-record management, and logout. It will also verify rejection of an incorrect master password and denial of record access after the session is locked.')

section('decimal'); chapter_positions.append((4,len(d.paragraphs)))
copy_range(237,282,skip=[257,277,281],joins={253:src.paragraphs[253].text+' '+src.paragraphs[257].text,276:src.paragraphs[276].text+' '+src.paragraphs[277].text+' '+src.paragraphs[281].text})
head('System Architecture',2)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
p.add_run().add_picture(str(ROOT/'architecture.png'),width=Inches(6.1))
p=para('Figure 4.1 Secure password manager system architecture','Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
para(src.paragraphs[309].text)
head('Failure Handling and Test Evidence')
para('An authentication failure should keep the vault locked and return a general error message without revealing whether protected records are present. A record that fails integrity verification should not be displayed as a partially recovered password. Storage errors should be reported without printing credentials, derived keys, or sensitive notes. These conditions will be included in the planned negative test cases.')
para('Each test record should identify the input conditions, action, expected behavior, observed behavior, and pass or fail decision. Screenshots used as evidence should contain synthetic data only. Timing observations should state the device, record count, and operation being measured so that results can be interpreted within the project scope.')

def table(title,rows,widths,description):
    para(title,'Caption')
    t=d.add_table(rows=0, cols=len(rows[0])); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for j,row in enumerate(rows):
        cells=t.add_row().cells
        for k,(c,text) in enumerate(zip(cells,row)):
            c.width=Inches(widths[k]); c.text=clean(text)
            c.vertical_alignment=1
            tcpr=c._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7E6E6' if j==0 else ('F6F6F6' if j%2==0 else 'FFFFFF')); tcpr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                a=OxmlElement('w:'+side); a.set(qn('w:w'),'85'); a.set(qn('w:type'),'dxa'); mar.append(a)
            tcpr.append(mar)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.1; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.keep_with_next=False
                for r in p.runs: r.font.size=Pt(11); r.bold=(j==0)
        trpr=t.rows[-1]._tr.get_or_add_trPr()
        if j==0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
        if j<len(rows)-1:
            for c in cells:
                for p in c.paragraphs: p.paragraph_format.keep_with_next=True
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    para(description)
    return t

head('Hardware Requirements')
rows=[[clean(c.text) for c in r.cells] for r in src.tables[0].rows]
rows[2][2]='Application, database, and test-data storage'; rows[4][2]='Package installation and optional updates'
table('Table 4.1 Hardware requirements',rows,[1.3,2.4,2.55],'Table 4.1 lists the development platform and supporting resources. The specifications are proposed laboratory requirements, rather than measured minimum requirements. Local vault operations are intended to remain available without a continuous internet connection.')
copy_range(322,337)

section('decimal'); chapter_positions.append((5,len(d.paragraphs)))
copy_range(368,380)
head('Expected Date of Completion of Work')
rows=[[clean(c.text) for c in r.cells] for r in src.tables[1].rows]
table('Table 5.1 Proposed activity schedule',rows,[4.2,2.05],'Table 5.1 presents the June-September 2026 schedule retained from the original project plan. The dates describe the planned sequence of work and do not establish actual completion. Requirements and design precede development, followed by validation and documentation.')
para('Progress reviews should confirm that each stage meets its acceptance checks before the next stage begins. A delay in the authentication or encryption workflow should be reflected in the integration and testing plan, rather than reducing the security checks.')
head('Approximate Expenditure')
rows=[[clean(c.text) for c in r.cells] for r in src.tables[2].rows]
for row in rows[1:-1]: row[2]='Not specified'
rows[-1]=['','Total estimated expenditure','','2,500']
table('Table 5.2 Estimated project expenditure',rows,[.55,3.05,1.15,1.5],'Table 5.2 summarizes the original estimated expenditure of Rs. 2,500. The figures are budget estimates and are not records of purchases. Software costs apply only if licenses are required; the final expenditure should be recorded from actual receipts.')
copy_range(394,412,skip=[403],joins={402:src.paragraphs[402].text+' '+src.paragraphs[403].text})
copy_range(412,415)

chapter(431,492,6)

chapter(492,498,7)
head('Planned Validation and Handover')
para('The next stage will translate the proposed modules into an integrated prototype using established cryptographic libraries. The team will first validate the master-password, salt, and key-derivation workflow, then test encrypted record persistence before connecting the credential-management interface. Password utilities and session controls will be added after the vault workflow behaves consistently. This order makes it easier to isolate failures in security-critical operations.')
head('Submission and Approval')
para('Date of Mini Project Phase I submission: __________________')
for name in ['Mr. Siraj Riyaj Sayyad | Roll No. 70','Mr. Prathamesh Jayantkumar Sathe | Roll No. 68','Mr. Adesh Dattatray Dethe | Roll No. 46']:
    para(name+' | Signature: __________________')
para('Project Guide: Ms. A. S. Pathan | Signature: __________________')
para('Head of Department: Dr. Vanita Jadhav | Signature: __________________')

section('decimal')
copy_range(518,550,skip=[524],joins={523:src.paragraphs[523].text+' '+src.paragraphs[524].text})

# Consistent main headings and searchable fields for automatic contents.
for p in d.paragraphs:
    if p.text=='Milestones': p.text='Planned Milestones'
    if p.style.name=='Heading 1': p.paragraph_format.space_before=Pt(0)
    if p.text.startswith('Chapter '): p.paragraph_format.keep_with_next=True
# Fit short chapters as complete pages, retaining comfortable academic typography.
current='front'
for p in d.paragraphs:
    if p.text=='Abstract': current='abstract'
    elif p.text=='Contents': current='contents'
    elif p.text.startswith('Chapter '): current=p.text.split(':')[0]
    elif p.text=='References': current='references'
    if current in ['abstract','Chapter 3','Chapter 5','Chapter 7'] and not p.style.name.startswith('Heading'):
        p.paragraph_format.line_spacing=1.12 if current not in ['Chapter 5','Chapter 7'] else 1.0
        p.paragraph_format.space_after=Pt(4)
        if current=='Chapter 5':
            p.paragraph_format.space_after=Pt(2)
            for r in p.runs: r.font.size=Pt(11.5)
    if current=='references' and not p.style.name.startswith('Heading'):
        p.style=d.styles['Normal']
        p.paragraph_format.line_spacing=1.0
        p.paragraph_format.space_after=Pt(3)
        for r in p.runs: r.font.size=Pt(10.5)
settings=d.settings.element
update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
d.core_properties.title='Secure Password Manager with Master-Key Encryption'
d.core_properties.subject='Mini Project Phase I Report'
d.save(OUT)
print(OUT)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))

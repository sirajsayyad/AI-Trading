from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import re
from docx import Document

ROOT=Path('C:/Users/siraj/Downloads/Secure Password Manager Overleaf')
SRC=Path('C:/Users/siraj/Downloads/Mini project Phase 1 report SS.docx')
doc=Document(SRC)
(ROOT/'images').mkdir(exist_ok=True)
(ROOT/'chapters').mkdir(exist_ok=True)
with ZipFile(SRC) as z:
    (ROOT/'images/college-logo.jpg').write_bytes(z.read('word/media/image1.jpeg'))

def clean(t):
    t=re.sub(r'(\d)\ufffd(?=\d)',r'\1–',t)
    t=t.replace('SVERI\ufffdS',"SVERI'S").replace('Flor\ufffdncio','Florêncio')
    t=re.sub(r'\ufffd(?=s\b)',"'",t)
    t=t.replace('users\ufffd',"users'").replace(' \ufffd ',' – ').replace('\ufffd','"')
    return ' '.join(t.split())

def tex(t):
    replacements={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(replacements.get(c,c) for c in clean(t))

def extract(a,b,skip=(),joins=None):
    lines=[]; listing=False; joins=joins or {}
    for i in range(a,b):
        if i in skip: continue
        p=doc.paragraphs[i]; t=clean(joins.get(i,p.text))
        if not t or t.isdigit() or t in ['storage','ates']: continue
        islist=p.style.name=='List Paragraph'
        if listing and not islist: lines.append('\\end{itemize}\n'); listing=False
        if islist:
            if not listing: lines.append('\\begin{itemize}'); listing=True
            lines.append('\\item '+tex(t))
        elif t.startswith('Chapter '): lines.append('\\chapter{'+tex(t.split(':',1)[1].strip())+'}\n')
        elif p.style.name.startswith('Heading'):
            level=int(p.style.name[-1]); cmd={1:'chapter',2:'section',3:'subsection'}.get(level,'subsubsection')
            lines.append('\\'+cmd+'{'+tex(t)+'}\n')
        else: lines.append(tex(t)+'\n')
    if listing: lines.append('\\end{itemize}\n')
    return '\n'.join(lines)

main=r'''% Upload this entire project ZIP to Overleaf and select XeLaTeX.
\documentclass[12pt,a4paper,oneside]{report}
\usepackage{fontspec}
% Optional: upload your own licensed Times New Roman files into fonts/.
\IfFileExists{fonts/times.ttf}{%
  \setmainfont{times.ttf}[Path=fonts/,BoldFont=timesbd.ttf,
    ItalicFont=timesi.ttf,BoldItalicFont=timesbi.ttf]
}{%
  \IfFontExistsTF{Times New Roman}{\setmainfont{Times New Roman}}%
    {\setmainfont{TeX Gyre Termes}}%
}
\usepackage[left=29mm,right=22mm,top=22mm,bottom=22mm]{geometry}
\usepackage{graphicx}
\usepackage{setspace}
\usepackage{array,tabularx,booktabs}
\usepackage{enumitem}
\usepackage{caption}
\usepackage{float}
\usepackage{placeins}
\usepackage{titlesec}
\usepackage{tikz}
\usetikzlibrary{arrows.meta,positioning}
\usepackage[hidelinks,unicode]{hyperref}
\hypersetup{pdftitle={Secure Password Manager with Master-Key Encryption},
  pdfauthor={Siraj Riyaj Sayyad, Prathamesh Jayantkumar Sathe, Adesh Dattatray Dethe}}
\setstretch{1.25}
\setlength{\parindent}{0pt}
\setlength{\parskip}{5pt}
\setlength{\emergencystretch}{2em}
\setlist[itemize]{leftmargin=1.5em,itemsep=3pt,topsep=4pt,parsep=0pt}
\captionsetup{font=small,labelfont=bf}
\captionsetup[table]{position=top,skip=6pt}
\renewcommand{\arraystretch}{1.2}
\titleformat{\chapter}[hang]{\normalfont\Large\bfseries}{Chapter \thechapter:}{0.6em}{}
\titlespacing*{\chapter}{0pt}{0pt}{16pt}
\titleformat{\section}{\normalfont\large\bfseries}{\thesection}{0.6em}{}
\titleformat{\subsection}{\normalfont\normalsize\bfseries}{\thesubsection}{0.6em}{}
\setcounter{tocdepth}{1}
\raggedbottom

% Current preference: no printed page numbers and no page borders.
% Change false to true below to print Roman preliminary page numbers
% and Arabic numbers beginning with Chapter 1.
\newif\ifshowpagenumbers
\showpagenumbersfalse
\ifshowpagenumbers
  \pagestyle{plain}
\else
  \pagestyle{empty}
  \makeatletter\let\ps@plain\ps@empty\makeatother
\fi

\begin{document}
\input{frontmatter}
\clearpage
\pagenumbering{arabic}
\input{chapters/chapter1}
\input{chapters/chapter2}
\input{chapters/chapter3}
\input{chapters/chapter4}
\input{chapters/chapter5}
\input{chapters/chapter6}
\input{chapters/chapter7}
\input{references}
\input{submission}
\end{document}
'''
(ROOT/'main.tex').write_text(main,encoding='utf-8')

front=r'''\begin{titlepage}
\centering
{\Large\bfseries Mini Project Phase I Report\par}
\vspace{8mm}
{\LARGE\bfseries Secure Password Manager with\\[2mm]Master-Key Encryption\par}
\vspace{8mm}
Submitted in partial fulfillment of the requirements for the award of\par
{\bfseries Final Year B.Tech. – Semester I\\Honors in Cyber Security\par}
\vspace{3mm}
{\bfseries CSE Hon-05B Mini Project\par}
\vspace{5mm}
Submitted to\par
{\bfseries Punyashlok Ahilyadevi Holkar Solapur University, Solapur\par}
\vspace{8mm}
{\bfseries Submitted by\par}
Mr. Siraj Riyaj Sayyad [Roll No. 70] [LY B]\\
Mr. Prathamesh Jayantkumar Sathe [Roll No. 68] [LY B]\\
Mr. Adesh Dattatray Dethe [Roll No. 46] [LY B]\par
{\bfseries Project Group No. B2\par}
\vspace{6mm}
Under the guidance of\par
{\bfseries Ms. A. S. Pathan\par}
\vspace{5mm}
\includegraphics[width=25mm]{images/college-logo.jpg}\par
\vfill
{\bfseries DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\par}
{\large\bfseries SVERI'S COLLEGE OF ENGINEERING, PANDHARPUR\par}
\vspace{3mm}
Academic Year 2026–2027
\end{titlepage}

\clearpage\thispagestyle{empty}
\begin{center}
\includegraphics[width=25mm]{images/college-logo.jpg}\par
\vspace{4mm}
{\large\bfseries SVERI'S COLLEGE OF ENGINEERING, PANDHARPUR\par}
{\small Affiliated to Punyashlok Ahilyadevi Holkar Solapur University, Solapur\par}
\vspace{10mm}
{\LARGE\bfseries CERTIFICATE\par}
\end{center}
\vspace{6mm}
This is to certify that the Mini Project Phase I Report entitled
\begin{center}
{\large\bfseries Secure Password Manager with Master-Key Encryption\par}
\end{center}
has been successfully submitted by
\begin{center}
Mr. Siraj Riyaj Sayyad [Roll No. 70] [LY B]\\
Mr. Prathamesh Jayantkumar Sathe [Roll No. 68] [LY B]\\
Mr. Adesh Dattatray Dethe [Roll No. 46] [LY B]
\end{center}
in partial fulfillment of the requirements for Final Year B.Tech. – Semester I
of Honors in Cyber Security for the subject CSE Hon-05B Mini Project,
as per the curriculum prescribed by Punyashlok Ahilyadevi Holkar Solapur
University, Solapur, during the Academic Year 2026–2027.

\vspace{15mm}
\noindent\begin{minipage}[t]{0.49\linewidth}
\textbf{Ms. A. S. Pathan}\\Project Guide
\end{minipage}\hfill
\begin{minipage}[t]{0.45\linewidth}
\textbf{Dr. Vanita Jadhav}\\Head of Department CSE
\end{minipage}
\vspace{15mm}
\begin{center}\textbf{Dr. B. P. Ronge}\\Principal\end{center}
\vspace{8mm}
Place: Pandharpur\hfill Date: \rule{32mm}{0.4pt}

\clearpage
\pagenumbering{roman}
\chapter*{Abstract}
\addcontentsline{toc}{chapter}{Abstract}
'''+tex(doc.paragraphs[134].text)+r'''

\clearpage
\tableofcontents
\clearpage
'''
(ROOT/'frontmatter.tex').write_text(front,encoding='utf-8')

(ROOT/'chapters/chapter1.tex').write_text(extract(130,133)+extract(135,169),encoding='utf-8')
(ROOT/'chapters/chapter2.tex').write_text(extract(169,218),encoding='utf-8')
(ROOT/'chapters/chapter3.tex').write_text(extract(218,237),encoding='utf-8')

architecture=r'''
\section{System Architecture}
\begin{figure}[H]
\centering
\begin{tikzpicture}[
  box/.style={draw,thick,align=center,text width=3.25cm,
    minimum height=1cm,inner sep=5pt,font=\small},
  side/.style={box,text width=2.65cm},
  flow/.style={-{Latex[length=2mm]},thick},
  node distance=8mm]
\node[box] (user) {User and application interface};
\node[box,below=of user] (auth) {Master-password authentication};
\node[box,below=of auth] (kdf) {Key derivation\\Salt and work factor};
\node[box,below=of kdf] (crypto) {Encryption and decryption};
\node[box,below=of crypto] (vault) {Encrypted credential vault};
\node[side,left=7mm of kdf] (manage) {Credential management\\Add, view, edit, delete, search};
\node[side,right=7mm of kdf] (utility) {Password generator\\Strength checker};
\node[side,left=7mm of vault] (session) {Session security\\Timeout, lock and logout};
\draw[flow] (user)--(auth);
\draw[flow] (auth)--(kdf);
\draw[flow] (kdf)--(crypto);
\draw[flow] (crypto)--(vault);
\draw[flow] (manage)--(kdf);
\draw[flow] (utility)--(kdf);
\draw[flow] (session.north)|-(crypto.west);
\end{tikzpicture}
\caption{Secure password manager system architecture}
\label{fig:architecture}
\end{figure}
'''

def table(caption,rows,colspec,label,desc):
    body=[]
    for i,row in enumerate(rows):
        cells=[tex(x) for x in row]
        if i==0: cells=[r'\textbf{'+x+'}' for x in cells]
        body.append(' & '.join(cells)+r' \\')
        if i==0: body.append(r'\midrule')
    return '\n'.join([r'\begin{table}[H]',r'\centering',r'\caption{'+caption+'}',r'\label{'+label+'}',r'\small',r'\begin{tabularx}{\linewidth}{'+colspec+'}',r'\toprule',*body,r'\bottomrule',r'\end{tabularx}',r'\end{table}',tex(desc),''])

rows=[[clean(c.text) for c in row.cells] for row in doc.tables[0].rows]
rows[2][2]='Application, database, and test-data storage'; rows[4][2]='Package installation and optional updates'
ch4=extract(237,282,skip=[257,277,281],joins={253:doc.paragraphs[253].text+' '+doc.paragraphs[257].text,276:doc.paragraphs[276].text+' '+doc.paragraphs[277].text+' '+doc.paragraphs[281].text})
ch4+=architecture+tex(doc.paragraphs[309].text)+'\n\n\\section{Hardware Requirements}\n'
ch4+=table('Hardware requirements',rows,r'@{}p{.21\linewidth}p{.35\linewidth}X@{}','tab:hardware','The hardware requirements describe the proposed development and testing platform. Storage supports the application and synthetic test records; internet access is needed for package installation and optional updates.')
ch4+=extract(322,337)
(ROOT/'chapters/chapter4.tex').write_text(ch4,encoding='utf-8')

ch5=extract(368,381)
rows=[[clean(c.text) for c in row.cells] for row in doc.tables[1].rows]
ch5+=table('Proposed activity schedule',rows,r'@{}X p{.30\linewidth}@{}','tab:schedule','The schedule records the June–September 2026 plan provided in the source report. It describes the intended sequence of activities and does not confirm actual completion.')
ch5+='\\section{Approximate Expenditure}\n'
rows=[[clean(c.text) for c in row.cells] for row in doc.tables[2].rows]
for row in rows[1:-1]: row[2]='--'
rows[-1]=['','Total','','2,500']
ch5+=table('Estimated project expenditure',rows,r'@{}p{.09\linewidth}X p{.14\linewidth}p{.19\linewidth}@{}','tab:budget','The estimated expenditure totals Rs. 2,500. These values are the original budget estimates and are not records of actual purchases. Quantities were not specified in the source report.')
ch5+=extract(394,412,skip=[403],joins={402:doc.paragraphs[402].text+' '+doc.paragraphs[403].text})
(ROOT/'chapters/chapter5.tex').write_text(ch5,encoding='utf-8')
(ROOT/'chapters/chapter6.tex').write_text(extract(431,492),encoding='utf-8')
(ROOT/'chapters/chapter7.tex').write_text(extract(492,498),encoding='utf-8')

refs=[]
for i in range(519,550):
    if i==524: continue
    t=doc.paragraphs[i].text
    if i==523: t+=' '+doc.paragraphs[524].text
    t=clean(t)
    if t and not t.isdigit(): refs.append(t)
bib=r'''\clearpage
\phantomsection
\addcontentsline{toc}{chapter}{References}
\renewcommand{\bibname}{References}
\begin{thebibliography}{99}
\small
\setlength{\itemsep}{4pt}
'''
for i,t in enumerate(refs,1): bib+='\\bibitem{ref'+str(i)+'} '+tex(t)+'\n\n'
bib+='\\end{thebibliography}\n'
(ROOT/'references.tex').write_text(bib,encoding='utf-8')

submission=r'''\FloatBarrier
\section*{Submission and Signatures}
Date of Mini Project Phase I submission: \rule{55mm}{0.4pt}

\vspace{6mm}
\textbf{Names, roll numbers and signatures of the students}

Mr. Siraj Riyaj Sayyad [Roll No. 70]\\Signature: \rule{65mm}{0.4pt}

Mr. Prathamesh Jayantkumar Sathe [Roll No. 68]\\Signature: \rule{65mm}{0.4pt}

Mr. Adesh Dattatray Dethe [Roll No. 46]\\Signature: \rule{65mm}{0.4pt}

\vspace{6mm}
Project Guide: Ms. A. S. Pathan\\Signature: \rule{65mm}{0.4pt}

\vspace{6mm}
Head of Department: Dr. Vanita Jadhav\\Signature: \rule{65mm}{0.4pt}
'''
(ROOT/'submission.tex').write_text(submission,encoding='utf-8')

readme='''# Secure Password Manager Mini Project Report

Upload the ZIP as a new Overleaf project. Set the compiler to XeLaTeX and the main document to main.tex, then Recompile. Compile twice if contents or references need another pass.

The project contains the seven chapters, cover, certificate, preliminary abstract, automatic contents, three tables, editable TikZ architecture figure, references, and submission signatures. Chapters automatically begin on new pages. No page border is used. Printed page numbers are disabled to follow your latest formatting request. To restore Roman preliminary numbers and Arabic chapter numbers, change \\showpagenumbersfalse to \\showpagenumberstrue in main.tex.

Font: the project uses Times New Roman when it is available, with TeX Gyre Termes (a Times-compatible font) as the Overleaf fallback. For exact Times New Roman on Overleaf, upload your own licensed times.ttf, timesbd.ttf, timesi.ttf, and timesbi.ttf into a fonts/ folder. The main file detects those files automatically. Proprietary fonts are not included in this ZIP.

Source: Mini project Phase 1 report SS.docx. The original body content is retained, broken paragraphs are joined, fixed page-number artifacts are removed, and the abstract is placed before Chapter 1. The architecture is redrawn in editable TikZ from the report's described modules. Table captions and descriptions are supplied.

The dates, budget, names and academic year are retained from the source. Bibliography entries are transcribed from the source rather than independently verified; several original entries omit full publication details. Review these details before academic submission.
'''
(ROOT/'README.md').write_text(readme,encoding='utf-8')

zipout=ROOT.with_suffix('.zip')
with ZipFile(zipout,'w',ZIP_DEFLATED) as z:
    for p in sorted(ROOT.rglob('*')):
        if p.is_file() and (p.suffix in ['.tex','.jpg','.md']): z.write(p,p.relative_to(ROOT).as_posix())
print(ROOT)
print(zipout)
print('References:',len(refs))

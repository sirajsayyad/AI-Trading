from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\siraj\Downloads\LY Capstone Project_Progress-I Report - AI Trading.docx")
OUT = Path(r"E:\AI_Trading\LY Capstone Project_Progress-I Report - AI Trading - Corrected.docx")

doc = Document(str(SRC))

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

def set_run_font(run, size=12, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def apply_para(para, *, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
               before=0, after=6, first_indent=0.3, line=1.5):
    fmt = para.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.alignment = align
    fmt.first_line_indent = Inches(first_indent) if first_indent else None
    for run in para.runs:
        set_run_font(run, size=size, bold=bold if bold else None)

def add_after(para, text):
    new = doc.add_paragraph(text)
    para._p.addnext(new._p)
    return new

# Split the very long methodology into readable paragraphs at sentence boundaries.
methodology = doc.paragraphs[99]
sentences = re.split(r"(?<=[.!?])\s+", methodology.text.strip())
chunks, current = [], []
for sentence in sentences:
    current.append(sentence)
    if len(current) >= 4:
        chunks.append(" ".join(current))
        current = []
if current:
    chunks.append(" ".join(current))
methodology.text = chunks[0]
anchor = methodology
for chunk in chunks[1:]:
    anchor = add_after(anchor, chunk)

# Re-fetch paragraphs after the split and style the entire report consistently.
section_heading_indices = set()
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if re.match(r"^(?:[1-9]|1[0-5])\.\s", text) or text == "References:":
        section_heading_indices.add(i)

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    if i in (1, 3, 21):
        apply_para(para, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=10, first_indent=0, line=1.15)
        continue
    if i in (2, 5, 6, 7, 8, 13, 14, 17, 18, 19, 22, 23, 24, 26):
        apply_para(para, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=6, first_indent=0, line=1.15)
        continue
    if i in section_heading_indices:
        apply_para(para, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, first_indent=0, line=1.15)
        continue
    if text.startswith("Figure ") or text.startswith("(Minimum") or text in {
        "Include:", "Design or model developed", "Demonstrations or prototype images",
        "Analysis, results, graphs, and discussion", "Attach photographs or screenshots with proper labels.",
        "At least 03 pages", "List remaining tasks with a tentative timeline or Gantt chart until project completion."
    }:
        apply_para(para, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, before=3, after=3, first_indent=0, line=1.15)
        continue
    apply_para(para, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, first_indent=0.3, line=1.5)

for i in (1, 3, 21, 23, 24, 26):
    if i < len(doc.paragraphs):
        doc.paragraphs[i].paragraph_format.first_line_indent = None

doc.save(str(OUT))
print(f"saved {OUT}")
